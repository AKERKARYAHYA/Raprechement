# -*- coding: utf-8 -*-
import streamlit as st
import pandas as pd
import re
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# CONFIG & STYLES (Clean & Professional)
# ─────────────────────────────────────────────
st.markdown("""
<style>
    .centered-header {
        text-align: center;
        color: #1e293b;
        font-size: 1.8rem;
        font-weight: 700;
        margin-top: 5px;
        margin-bottom: 25px;
    }
    .metric-card {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-left: 4px solid #1e3a8a;
        border-radius: 8px;
        padding: 0.8rem 1rem;
        box-shadow: 0 1px 2px rgba(0,0,0,0.05);
    }
    .metric-value {
        font-size: 1.6rem;
        font-weight: 700;
        color: #0f172a;
    }
    .metric-label {
        font-size: 0.75rem;
        color: #64748b;
        text-transform: uppercase;
        letter-spacing: 0.05em;
    }
    .success-box {
        background: #f0fdf4;
        border-left: 4px solid #16a34a;
        border-radius: 6px;
        padding: 0.8rem 1rem;
        color: #15803d;
        margin: 0.5rem 0;
    }
    .warning-box {
        background: #fffbe3;
        border-left: 4px solid #d97706;
        border-radius: 6px;
        padding: 0.8rem 1rem;
        color: #b45309;
        margin: 0.5rem 0;
    }
    .section-header {
        font-size: 1.1rem;
        font-weight: 600;
        color: #1e293b;
        border-bottom: 2px solid #e2e8f0;
        padding-bottom: 0.3rem;
        margin: 1.2rem 0 0.8rem 0;
    }
    /* Izalat l-khotot l-motaqatti3a (dashed borders) dyal uploader */
    div[data-testid="stFileUploader"] {
        border: none !important;
        padding: 0rem;
    }
    div[data-testid="stFileUploader"] > section {
        border: none !important;
    }
    /* Style dyal l-bouton zrayq (Bleu Roi #1e3a8a) */
    div.stButton > button {
        background-color: #1e3a8a !important;
        color: #ffffff !important;
        border: none !important;
        border-radius: 6px !important;
        font-weight: 600 !important;
    }
    div.stButton > button:hover {
        background-color: #1d4ed8 !important;
        color: #ffffff !important;
    }
</style>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────

def normalize_cols(df: pd.DataFrame) -> pd.DataFrame:
    """Lowercase + strip all column names."""
    df.columns = [str(c).strip().lower() for c in df.columns]
    return df

def find_col(df: pd.DataFrame, candidates: list) -> str | None:
    """Return first matching column name (case-insensitive)."""
    for c in candidates:
        if c.lower() in df.columns:
            return c.lower()
    return None

def extract_date_from_description(desc: str) -> str | None:
    """Extract DD/MM/YYYY from a description string."""
    if not isinstance(desc, str):
        return None
    m = re.search(r'\b(\d{2}/\d{2}/\d{4})\b', desc)
    return m.group(1) if m else None

def extract_id_from_erp_desc(desc: str) -> str | None:
    """Extract numeric ID at the end of ERP description."""
    if not isinstance(desc, str):
        return None
    m = re.search(r'(\d+)\s*$', desc.strip())
    return m.group(1) if m else None

def detect_city(desc: str, city_keywords: dict) -> str | None:
    """Detect RLV city from description using keyword mapping."""
    if not isinstance(desc, str):
        return None
    desc_up = desc.upper()
    for keyword, city in city_keywords.items():
        if keyword.upper() in desc_up:
            return city
    return None

CITY_KEYWORDS = {
    "FES": "Fes",
    "FÈS": "Fes",
    "MARRAKECH": "Marrakech",
    "MARRAKESH": "Marrakech",
    "RABAT": "Rabat",
    "AGADIR": "Agadir",
}

GEO_MAP = {
    "Fes":        ["Meknes"],
    "Marrakech": ["Marrakech", "Essaouira"],
    "Rabat":      ["Rabat"],
    "Agadir":     ["Agadir"],
}

def detect_erp_city(desc: str, city_keywords: dict) -> str | None:
    """Detect ERP city from description."""
    if not isinstance(desc, str):
        return None
    desc_up = desc.upper()
    erp_cities = ["MEKNES", "MÉKNÈS", "MARRAKECH", "MARRAKESH", "ESSAOUIRA", "RABAT", "AGADIR"]
    for kw in erp_cities:
        if kw in desc_up:
            canonical = {
                "MEKNES": "Meknes", "MÉKNÈS": "Meknes",
                "MARRAKECH": "Marrakech", "MARRAKESH": "Marrakech",
                "ESSAOUIRA": "Essaouira",
                "RABAT": "Rabat",
                "AGADIR": "Agadir",
            }
            return canonical.get(kw)
    return None

def style_excel_sheet(ws, header_fill_hex: str, col_widths: list):
    """Apply professional formatting to an openpyxl worksheet."""
    header_fill = PatternFill("solid", start_color=header_fill_hex, end_color=header_fill_hex)
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="C0C8D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, cell in enumerate(ws[1], start=1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border
        if col_idx <= len(col_widths):
            ws.column_dimensions[get_column_letter(col_idx)].width = col_widths[col_idx - 1]

    alt_fill = PatternFill("solid", start_color="F4F7FA", end_color="F4F7FA")
    data_font = Font(name="Arial", size=10)
    data_align = Alignment(vertical="center")

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.font = data_font
            cell.alignment = data_align
            cell.border = border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    ws.row_dimensions[1].height = 35
    ws.freeze_panes = "A2"

def build_excel(matched_df: pd.DataFrame, unmatched_df: pd.DataFrame) -> bytes:
    """Create styled Excel workbook and return as bytes."""
    wb = Workbook()

    # Sheet 1: Matched_Results
    ws1 = wb.active
    ws1.title = "Matched_Results"

    matched_cols = [
        "reference", "description", "nom", "n°document", "debit", "credit",
        "la date", "description_bank", "debit_bank", "credit_bank",
        "Ecart", "date RLV"
    ]
    present_matched = [c for c in matched_cols if c in matched_df.columns]
    ws1.append(present_matched)
    for _, row in matched_df[present_matched].iterrows():
        ws1.append([row[c] for c in present_matched])

    col_widths_matched = [15, 35, 20, 15, 14, 14, 14, 35, 14, 14, 12, 14]
    style_excel_sheet(ws1, "1A3A5C", col_widths_matched[:len(present_matched)])

    ecart_col_idx = present_matched.index("Ecart") + 1 if "Ecart" in present_matched else None
    if ecart_col_idx:
        green_fill = PatternFill("solid", start_color="C8E6C9", end_color="C8E6C9")
        orange_fill = PatternFill("solid", start_color="FFE0B2", end_color="FFE0B2")
        for row_idx in range(2, ws1.max_row + 1):
            cell = ws1.cell(row=row_idx, column=ecart_col_idx)
            try:
                val = float(cell.value or 0)
                cell.fill = green_fill if val == 0 else orange_fill
            except (TypeError, ValueError):
                pass

    # Sheet 2: Unmatched_RLV
    ws2 = wb.create_sheet("Unmatched_RLV")
    unmatched_cols = ["la date", "description", "debit", "credit"]
    present_unmatched = [c for c in unmatched_cols if c in unmatched_df.columns]
    ws2.append(present_unmatched)
    for _, row in unmatched_df[present_unmatched].iterrows():
        ws2.append([row[c] for c in present_unmatched])

    style_excel_sheet(ws2, "B71C1C", [14, 50, 14, 14])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

# ─────────────────────────────────────────────
# CORE RECONCILIATION
# ─────────────────────────────────────────────

def reconcile(erp_df: pd.DataFrame, rlv_df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    logs = []

    erp_df = normalize_cols(erp_df.copy())
    rlv_df = normalize_cols(rlv_df.copy())

    erp_credit_col = find_col(erp_df, ["credit", "crédit", "montant_credit"])
    erp_debit_col  = find_col(erp_df, ["debit", "débit", "montant_debit"])
    erp_date_col   = find_col(erp_df, ["date", "date_ecriture", "date_comptable"])
    erp_desc_col   = find_col(erp_df, ["description", "libelle", "libellé", "description_erp"])
    erp_ref_col    = find_col(erp_df, ["reference", "référence", "ref"])
    erp_nom_col    = find_col(erp_df, ["nom", "name"])
    erp_doc_col    = find_col(erp_df, ["n°document", "n° document", "ndocument", "document", "num_doc", "numero_document"])

    missing_erp = [k for k, v in {
        "credit": erp_credit_col, "date": erp_date_col, "description": erp_desc_col
    }.items() if v is None]
    if missing_erp:
        st.error(f"Fichier ERP manquant de colonnes requises: {missing_erp}")
        st.stop()

    rlv_credit_col = find_col(rlv_df, ["credit", "crédit", "montant_credit"])
    rlv_debit_col  = find_col(rlv_df, ["debit", "débit"])
    rlv_desc_col   = find_col(rlv_df, ["description", "libelle", "libellé", "description_bank", "description banque"])
    rlv_date_col   = find_col(rlv_df, ["la date", "date", "date_valeur", "date valeur"])

    missing_rlv = [k for k, v in {
        "credit": rlv_credit_col, "description": rlv_desc_col
    }.items() if v is None]
    if missing_rlv:
        st.error(f"Fichier RLV manquant de colonnes requises: {missing_rlv}")
        st.stop()

    erp_df[erp_credit_col] = pd.to_numeric(erp_df[erp_credit_col], errors="coerce").fillna(0)
    rlv_df[rlv_credit_col] = pd.to_numeric(rlv_df[rlv_credit_col], errors="coerce").fillna(0)

    erp_credit = erp_df[erp_df[erp_credit_col] > 0].copy().reset_index(drop=True)
    rlv_credit = rlv_df[rlv_df[rlv_credit_col] > 0].copy().reset_index(drop=True)

    logs.append(f"Lignes Crédit ERP: {len(erp_credit)} / Lignes Crédit RLV: {len(rlv_credit)}")

    rlv_credit["_rlv_date"] = rlv_credit[rlv_desc_col].apply(extract_date_from_description)
    erp_credit["_erp_id"]   = erp_credit[erp_desc_col].apply(extract_id_from_erp_desc)
    erp_credit["_erp_city"] = erp_credit[erp_desc_col].apply(lambda x: detect_erp_city(x, {}))

    erp_credit["_erp_date_str"] = pd.to_datetime(
        erp_credit[erp_date_col], dayfirst=True, errors="coerce"
    ).dt.strftime("%d/%m/%Y")

    rlv_credit["_rlv_city"] = rlv_credit[rlv_desc_col].apply(
        lambda x: detect_city(x, CITY_KEYWORDS)
    )

    grp_cols = ["_erp_id", "_erp_city", "_erp_date_str"]
    erp_grouped = (
        erp_credit.groupby(grp_cols, dropna=False)
        .agg(
            _sum_credit=(erp_credit_col, "sum"),
            _row_indices=("_erp_id", lambda x: list(x.index)),
        )
        .reset_index()
    )

    matched_rows    = []
    matched_rlv_idx = set()

    for rlv_idx, rlv_row in rlv_credit.iterrows():
        rlv_city   = rlv_row["_rlv_city"]
        rlv_date   = rlv_row["_rlv_date"]
        rlv_amount = float(rlv_row[rlv_credit_col])

        if rlv_city is None or rlv_date is None:
            continue

        erp_targets = GEO_MAP.get(rlv_city, [])
        matched = False

        for erp_city_target in erp_targets:
            candidates = erp_grouped[
                (erp_grouped["_erp_city"] == erp_city_target) &
                (erp_grouped["_erp_date_str"] == rlv_date)
            ]

            for _, grp_row in candidates.iterrows():
                ecart = rlv_amount - grp_row["_sum_credit"]
                if -700 <= ecart <= 700:
                    for erp_row_idx in grp_row["_row_indices"]:
                        erp_r = erp_credit.loc[erp_row_idx]
                        row_out = {
                            "reference":        erp_r.get(erp_ref_col, "") if erp_ref_col else "",
                            "description":      erp_r.get(erp_desc_col, ""),
                            "nom":              erp_r.get(erp_nom_col, "") if erp_nom_col else "",
                            "n°document":        erp_r.get(erp_doc_col, "") if erp_doc_col else "",
                            "debit":            erp_r.get(erp_debit_col, 0) if erp_debit_col else 0,
                            "credit":           erp_r.get(erp_credit_col, 0),
                            "la date":          rlv_date,
                            "description_bank": rlv_row[rlv_desc_col],
                            "debit_bank":       rlv_row.get(rlv_debit_col, 0) if rlv_debit_col else 0,
                            "credit_bank":      rlv_amount,
                            "Ecart":            round(ecart, 2),
                            "date RLV":         rlv_row[rlv_date_col] if rlv_date_col else "",
                        }
                        matched_rows.append(row_out)

                    matched_rlv_idx.add(rlv_idx)
                    erp_grouped = erp_grouped.drop(grp_row.name)
                    matched = True
                    break

            if matched:
                break

    unmatched_rlv = rlv_credit[~rlv_credit.index.isin(matched_rlv_idx)].copy()

    matched_df = pd.DataFrame(matched_rows) if matched_rows else pd.DataFrame(
        columns=["reference","description","nom","n°document","debit","credit",
                 "la date","description_bank","debit_bank","credit_bank","Ecart","date RLV"]
    )

    unmatched_out_cols = {}
    if rlv_date_col:
        unmatched_out_cols["la date"] = unmatched_rlv[rlv_date_col]
    else:
        unmatched_out_cols["la date"] = unmatched_rlv["_rlv_date"]
    unmatched_out_cols["description"] = unmatched_rlv[rlv_desc_col]
    unmatched_out_cols["debit"]       = unmatched_rlv[rlv_debit_col] if rlv_debit_col else 0
    unmatched_out_cols["credit"]      = unmatched_rlv[rlv_credit_col]
    unmatched_df = pd.DataFrame(unmatched_out_cols).reset_index(drop=True)

    stats = {
        "erp_total":      len(erp_credit),
        "rlv_total":      len(rlv_credit),
        "matched":        len(matched_rlv_idx),
        "unmatched":      len(unmatched_df),
        "match_rate":      round(len(matched_rlv_idx) / max(len(rlv_credit), 1) * 100, 1),
        "logs":           logs,
    }

    return matched_df, unmatched_df, stats

# ─────────────────────────────────────────────
# UI PRINCIPALE
# ─────────────────────────────────────────────

st.markdown('<div class="centered-header">Rapprochement Espèces — CAML</div>', unsafe_allow_html=True)

# ── File Upload 
col1, col2 = st.columns(2)
with col1:
    st.markdown('<div class="section-header">Fichier ERP</div>', unsafe_allow_html=True)
    erp_file = st.file_uploader("Importer le fichier ERP (Excel/CSV)", type=["xlsx", "xls", "csv"], key="erp_caml")

with col2:
    st.markdown('<div class="section-header">Relevé bancaire (RLV)</div>', unsafe_allow_html=True)
    rlv_file = st.file_uploader("Importer le Relevé Bancaire (Excel/CSV)", type=["xlsx", "xls", "csv"], key="rlv_caml")

st.markdown("")
run_col = st.columns([1, 2, 1])[1]
with run_col:
    run_btn = st.button("🚀 Lancer le rapprochement des données", use_container_width=True)

if run_btn:
    if erp_file and rlv_file:
        st.markdown('<div class="section-header">Configuration des Fichiers</div>', unsafe_allow_html=True)
        cfg_col1, cfg_col2 = st.columns(2)

        def load_preview(file, label: str):
            ext = file.name.rsplit(".", 1)[-1].lower()
            if ext == "csv":
                return None, None, pd.read_csv(file, nrows=5)
            xf = pd.ExcelFile(file)
            sheets = xf.sheet_names
            sheet = st.selectbox(f"{label} — Feuille", sheets, key=f"{label}_sheet")
            header_row = st.number_input(f"{label} — Ligne d'en-tête (1-based)", min_value=1, max_value=20, value=1, key=f"{label}_hdr")
            df_prev = pd.read_excel(file, sheet_name=sheet, header=header_row - 1, nrows=5)
            return sheet, header_row, df_prev

        with cfg_col1:
            erp_sheet, erp_hdr, erp_prev = load_preview(erp_file, "ERP")
            st.dataframe(erp_prev, use_container_width=True)

        with cfg_col2:
            rlv_sheet, rlv_hdr, rlv_prev = load_preview(rlv_file, "RLV")
            st.dataframe(rlv_prev, use_container_width=True)

        with st.spinner("Traitement en cours..."):
            def read_file(file, sheet, hdr):
                ext = file.name.rsplit(".", 1)[-1].lower()
                file.seek(0)
                if ext == "csv":
                    return pd.read_csv(file)
                return pd.read_excel(file, sheet_name=sheet, header=hdr - 1)

            erp_df = read_file(erp_file, erp_sheet, erp_hdr)
            rlv_df = read_file(rlv_file, rlv_sheet, rlv_hdr)

            matched_df, unmatched_df, stats = reconcile(erp_df, rlv_df)
            excel_bytes = build_excel(matched_df, unmatched_df)

        st.markdown('<div class="section-header">Résultats</div>', unsafe_allow_html=True)
        k1, k2, k3, k4, k5 = st.columns(5)

        def kpi(col, label, value, color="#1e3a8a"):
            col.markdown(f"""
            <div class="metric-card" style="border-left-color:{color}">
                <div class="metric-label">{label}</div>
                <div class="metric-value" style="color:{color}">{value}</div>
            </div>""", unsafe_allow_html=True)

        kpi(k1, "Lignes ERP",       stats["erp_total"])
        kpi(k2, "Lignes RLV",       stats["rlv_total"])
        kpi(k3, "Rapprochées",     stats["matched"],   "#16a34a")
        kpi(k4, "Non Rapprochées", stats["unmatched"], "#dc2626")
        kpi(k5, "Taux",             f"{stats['match_rate']}%",
            "#16a34a" if stats["match_rate"] >= 70 else "#d97706")

        st.markdown("")
        dl_col = st.columns([1, 2, 1])[1]
        with dl_col:
            st.download_button(
                label="Télécharger le Rapport (.xlsx)",
                data=excel_bytes,
                file_name="CAML_Rapprochement_Results.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

        st.markdown('<div class="section-header">Aperçu des Données</div>', unsafe_allow_html=True)
        tab1, tab2 = st.tabs([
            f"Rapprochées ({len(matched_df)} lignes)",
            f"Non Rapprochées RLV ({len(unmatched_df)} lignes)"
        ])

        with tab1:
            if matched_df.empty:
                st.markdown('<div class="warning-box">Aucune ligne rapprochée trouvée.</div>', unsafe_allow_html=True)
            else:
                st.dataframe(matched_df, use_container_width=True, height=400)

        with tab2:
            if unmatched_df.empty:
                st.markdown('<div class="success-box">Toutes les lignes RLV ont été rapprochées !</div>', unsafe_allow_html=True)
            else:
                st.dataframe(unmatched_df, use_container_width=True, height=400)

    else:
        st.warning("⚠️ Veuillez d'abord importer les deux fichiers ERP et RLV !")