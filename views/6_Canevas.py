"""
Canevas ERP Generator - Correction Erreur Duplicate Columns
"""

import io
import re
import unicodedata
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Configuration des colonnes ERP
# --------------------------------------------------------------------------- #
ERP_COLUMNS = [
    "JOURNALNUM", "LINENUM", "VOUCHERSTD", "TRDATEDMY", "ACCOUNTTYPE",
    "CUSTACCOUNT", "DEBIT", "CREDIT", "DEVISE", "EXCHANGERATE", "NUMDOCUMENT",
    "DATEDOCUMENT", "PROFILEVALIDATION", "ACTIVITE", "AGENCEAXE", "BU",
    "COSTCENTER", "INTERCOS", "CODETAXE", "GROUPTAXE", "INVOICEID", "EXCHANGERATE_2", "APPROVER",
    "APPROVED", "PAYMODE", "TYPE_COMPTE", "COMPTE", "Description",
    "BENEFICIEDID", "BENEFICIERBANK", "REFpay",
]

TEXT_COLUMNS = ["ACCOUNTTYPE", "CUSTACCOUNT", "COMPTE", "TRDATEDMY", "ACTIVITE", "Description", "AGENCEAXE"]

# --------------------------------------------------------------------------- #
# Fonctions de nettoyage
# --------------------------------------------------------------------------- #
def normalize(text) -> str:
    s = unicodedata.normalize("NFKD", str(text))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


def is_blank(v) -> bool:
    if v is None:
        return True
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(v, str) and v.strip() == ""


def clean_account(v):
    if is_blank(v):
        return None
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def to_number(v):
    if is_blank(v):
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\u00a0", "").replace(" ", "")
    s = re.sub(r"[^\d,.\-+]", "", s)
    if not s:
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def derive_activity(description, kw_lait, kw_vro, priority, fallback):
    if is_blank(description):
        return fallback
    text = normalize(description)
    has_lait = bool(kw_lait) and normalize(kw_lait) in text
    has_vro = bool(kw_vro) and normalize(kw_vro) in text
    if has_lait and has_vro:
        return f"COPG_{priority}"
    if has_lait:
        return "COPG_LAIT"
    if has_vro:
        return "COPG_VRO"
    return fallback


# --------------------------------------------------------------------------- #
# Lecture et Auto-détection du fichier source
# --------------------------------------------------------------------------- #
@st.cache_data(show_spinner=False)
def list_sheets(file_bytes: bytes) -> list:
    return pd.ExcelFile(io.BytesIO(file_bytes)).sheet_names


@st.cache_data(show_spinner=False)
def read_raw(file_bytes: bytes, sheet: str) -> pd.DataFrame:
    return pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet, header=None, dtype=object)


def detect_header_row(raw: pd.DataFrame, scan: int = 20) -> int:
    best_idx, best_score = 0, -1
    for i in range(min(scan, len(raw))):
        score = sum(isinstance(x, str) and x.strip() != "" for x in raw.iloc[i].tolist())
        if score > best_score:
            best_idx, best_score = i, score
    return best_idx


def apply_header(raw: pd.DataFrame, header_idx: int) -> pd.DataFrame:
    names, seen = [], {}
    for i, h in enumerate(raw.iloc[header_idx].tolist()):
        name = str(h).strip() if not is_blank(h) else f"Colonne_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        names.append(name)
    df = raw.iloc[header_idx + 1:].copy()
    df.columns = names
    return df.dropna(how="all").reset_index(drop=True)


def auto_map_columns(columns: list) -> dict:
    mapping = {}
    for col in columns:
        n = normalize(col)
        if "type" in n and "compte" in n and "contrepartie" not in n:
            mapping["type"] = col
        elif n.startswith("compte") and "contrepartie" not in n and "type" not in n:
            mapping["compte"] = col
        elif "debit" in n:
            mapping["debit"] = col
        elif "contrepartie" in n:
            mapping["contrepartie"] = col
        elif "description" in n or "libelle" in n:
            mapping["description"] = col
        elif "agence" in n or "axe" in n:
            mapping["agence"] = col
    return mapping


# --------------------------------------------------------------------------- #
# Construction du Canevas ERP & Statistiques
# --------------------------------------------------------------------------- #
def build_canevas(src: pd.DataFrame, mapping: dict, rules: dict):
    n = len(src)

    debits = [to_number(v) for v in src[mapping["debit"]]] if mapping.get("debit") else [None] * n
    descriptions = [None if is_blank(v) else str(v).strip() for v in src[mapping["description"]]] if mapping.get("description") else [None] * n
    activities = [
        derive_activity(d, rules["kw_lait"], rules["kw_vro"], rules["priority"], rules["fallback"])
        for d in descriptions
    ]

    agences_for_stats = [None] * n
    if mapping.get("agence") and mapping["agence"] in src.columns:
        agences_for_stats = [None if is_blank(v) else str(v).strip() for v in src[mapping["agence"]]]

    out = pd.DataFrame({c: [None] * n for c in ERP_COLUMNS}, dtype=object)
    out["JOURNALNUM"] = None
    out["LINENUM"] = 1
    out["VOUCHERSTD"] = list(range(1, n + 1))
    
    if mapping.get("type") and mapping["type"] in src.columns:
        out["ACCOUNTTYPE"] = [None if is_blank(v) else str(v).strip() for v in src[mapping["type"]]]
    if mapping.get("compte") and mapping["compte"] in src.columns:
        out["CUSTACCOUNT"] = [clean_account(v) for v in src[mapping["compte"]]]
    
    out["DEBIT"] = debits
    out["CREDIT"] = None
    out["DEVISE"] = "MAD"
    
    # 🎯 EXCHANGERATE LAWLA FIHA 1
    out["EXCHANGERATE"] = 1
    
    out["NUMDOCUMENT"] = None
    out["DATEDOCUMENT"] = None
    out["PROFILEVALIDATION"] = "STD"
    out["ACTIVITE"] = activities
    out["AGENCEAXE"] = None
    out["BU"] = None
    out["COSTCENTER"] = None
    out["INTERCOS"] = None
    out["CODETAXE"] = None
    out["GROUPTAXE"] = None
    out["INVOICEID"] = None
    
    # 🎯 EXCHANGERATE TANYA (Khawya)
    out["EXCHANGERATE_2"] = None
    
    out["APPROVER"] = None
    out["APPROVED"] = "oui"
    out["PAYMODE"] = None
    out["TYPE_COMPTE"] = "Comptabilité"
    
    if mapping.get("contrepartie") and mapping["contrepartie"] in src.columns:
        out["COMPTE"] = [None if is_blank(v) else str(v).strip() for v in src[mapping["contrepartie"]] if not is_blank(v)]
    else:
        out["COMPTE"] = None
        
    out["Description"] = descriptions
    out["BENEFICIEDID"] = None
    out["BENEFICIERBANK"] = None
    out["REFpay"] = None

    # Copie dyal dataframe bach n-baynoha f l-interface b smiyat mo5talfin (EXCHANGERATE & EXCHANGERATE_2)
    display_df = out[ERP_COLUMNS].copy()

    # DataFrame liradi yt-généra f Excel (jouj b smiyta EXCHANGERATE)
    export_df = out[ERP_COLUMNS].copy()
    export_df.rename(columns={"EXCHANGERATE_2": "EXCHANGERATE"}, inplace=True)

    valid_rows = []
    for d, a in zip(debits, agences_for_stats):
        if d is not None:
            valid_rows.append((d, a if a is not None else "N/A"))

    if valid_rows:
        total_debit = float(sum(d for d, a in valid_rows))
        count_ops = len(valid_rows)
        max_d, max_ag = max(valid_rows, key=lambda x: x[0])
        min_d, min_ag = min(valid_rows, key=lambda x: x[0])
    else:
        total_debit = 0.0
        count_ops = 0
        max_d, max_ag = 0.0, "N/A"
        min_d, min_ag = 0.0, "N/A"

    stats = {
        "rows": n,
        "total_debit": total_debit,
        "count_ops": count_ops,
        "max_debit": max_d,
        "max_agence": max_ag,
        "min_debit": min_d,
        "min_agence": min_ag,
    }
    return display_df, export_df, stats


def to_excel_bytes(df: pd.DataFrame) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Canevas ERP", index=False)
        ws = writer.sheets["Canevas ERP"]

        base_font = Font(name="Arial", size=10)
        head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", start_color="1F4E78")
        col_index = {name: i + 1 for i, name in enumerate(df.columns)}

        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.font = base_font
        for name in TEXT_COLUMNS:
            for r in range(2, ws.max_row + 1):
                if name in col_index:
                    ws.cell(row=r, column=col_index[name]).number_format = "@"
        for r in range(2, ws.max_row + 1):
            if "DEBIT" in col_index:
                ws.cell(row=r, column=col_index["DEBIT"]).number_format = "#,##0.00"

        for name, idx in col_index.items():
            values = [len(str(v)) for v in df[name].dropna().head(500)] or [0]
            ws.column_dimensions[get_column_letter(idx)].width = min(max(len(name), max(values)) + 3, 50)
        ws.freeze_panes = "A2"
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Interface Streamlit
# --------------------------------------------------------------------------- #
def main():
    st.set_page_config(page_title="Générateur Canevas ERP", page_icon="📄", layout="wide")
    st.title("📄 Générateur de Canevas ERP & Tableau de Bord")
    st.caption("Importez votre fichier : EXCHANGERATE double gérée avec succès.")

    with st.sidebar:
        st.header("⚙️ Règle ACTIVITE")
        kw_lait = st.text_input("Mot-clé → LAIT", value="LAIT")
        kw_vro = st.text_input("Mot-clé → VRO", value="VRO")
        priority = st.radio("Si les deux mots-clés sont présents", ["LAIT", "VRO"], horizontal=True)
        fallback_label = st.selectbox("Si aucun mot-clé n'est trouvé", ["Laisser vide", "COPG_LAIT", "COPG_VRO"])
        fallback = None if fallback_label == "Laisser vide" else fallback_label

    st.subheader("1️⃣ Fichier source")
    uploaded = st.file_uploader("Fichier Excel (.xlsx ou .xls)", type=["xlsx", "xls"])
    if uploaded is None:
        st.info("Veuillez charger un fichier Excel pour commencer.")
        return

    file_bytes = uploaded.getvalue()
    try:
        sheets = list_sheets(file_bytes)
    except Exception as exc:
        st.error(f"Erreur : {exc}")
        return

    sheet = st.selectbox("Feuille à traiter", sheets) if len(sheets) > 1 else sheets[0]
    raw = read_raw(file_bytes, sheet)
    if raw.empty:
        st.warning("Feuille vide.")
        return

    default_header = detect_header_row(raw)
    src = apply_header(raw, default_header)
    if src.empty:
        st.warning("Aucune donnée.")
        return
    columns = list(src.columns)

    mapping = auto_map_columns(columns)

    if "contrepartie" not in mapping:
        st.error("⚠️ لم يتم العثور تلقائياً على عمود 'Compte de contrepartie'. تأكد من أن اسم العمود في ملف الإكسل يحتوي على كلمة contrepartie.")
        return

    st.subheader("2️⃣ Tableau de Bord & Statistiques Financières 📊")
    rules = {"kw_lait": kw_lait, "kw_vro": kw_vro, "priority": priority, "fallback": fallback}
    try:
        display_df, export_df, stats = build_canevas(src, mapping, rules)
    except Exception as exc:
        st.error(f"Erreur lors de la génération : {exc}")
        return

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("💰 Total Versements", f"{stats['total_debit']:,.2f} MAD")
    m2.metric("🔢 Nombre d'Opérations", f"{stats['count_ops']}")
    m3.metric("📈 Akbar Montant", f"{stats['max_debit']:,.2f} MAD", f"Agence: {stats['max_agence']}")
    m4.metric("📉 Asrar Montant", f"{stats['min_debit']:,.2f} MAD", f"Agence: {stats['min_agence']}")

    st.subheader("3️⃣ Aperçu du Canevas ERP")
    st.dataframe(display_df.head(200), use_container_width=True, hide_index=True)

    st.subheader("4️⃣ Export")
    xlsx_bytes = to_excel_bytes(export_df)
    stem = re.sub(r"\.xlsx?$", "", uploaded.name, flags=re.I)
    st.download_button(
        "⬇️ Télécharger le Canevas ERP (.xlsx)",
        data=xlsx_bytes,
        file_name=f"Canevas_ERP_Clean_{stem}_{datetime.now():%Y%m%d_%H%M}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
    )


if __name__ == "__main__":
    main()